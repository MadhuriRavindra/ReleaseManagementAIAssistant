"""
Builds the files attached to each change ticket: the release note (.docx) and the
TR list (.xlsx), like the attachments a release manager opens from ServiceNow.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NAVY = RGBColor(0x1F, 0x2A, 0x44)
GREY = RGBColor(0x6B, 0x72, 0x80)


def _kv_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for k, v in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = k, v
        cells[0].paragraphs[0].runs[0].bold = True
    return table


def release_note_docx(c, release_name, window, path):
    rn = c["release_note"]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10.5)

    head = doc.add_paragraph()
    run = head.add_run("Orbitra Retail · Change & Release Management")
    run.font.size, run.font.color.rgb = Pt(9), GREY
    title = doc.add_heading("Release Note", level=0)
    title.runs[0].font.color.rgb = NAVY
    sub = doc.add_paragraph()
    r = sub.add_run(f"{c['change_id']} · {c['short_description']}")
    r.bold, r.font.size = True, Pt(12)

    doc.add_heading("Change details", level=1)
    _kv_table(doc, [
        ("Change", c["change_id"]),
        ("Application (CI)", c["application"]),
        ("System", c["system"]),
        ("Release", f"{c['requested_release']} ({release_name})"),
        ("Planned window", window),
        ("Category / type", f"{c['work_type']} / {c['change_type']}"),
        ("Risk / priority", f"{c['risk']} / {c['priority']}"),
        ("Downtime", rn["downtime"]),
        ("Assignment group", c["assignment_group"]),
    ])

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(rn["summary"])

    doc.add_heading("Deployment steps", level=1)
    for step in rn["deployment_steps"]:
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("Transport requests", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for cell, label in zip(table.rows[0].cells, ["#", "TR", "Type", "A / P target",
                                                  "Import after"]):
        cell.text = label
    for t in c["tr_list"]:
        cells = table.add_row().cells
        cells[0].text = str(t["sequence"])
        cells[1].text = t["tr_id"]
        cells[2].text = t["type"]
        cells[3].text = f"{t['a_target']} / {t['p_target']}"
        cells[4].text = t["depends_on_tr"] or "—"

    doc.add_heading("Dependencies", level=1)
    doc.add_paragraph(rn["dependencies"])
    doc.add_heading("Rollback", level=1)
    doc.add_paragraph(rn["rollback"])

    doc.add_heading("Sign-off", level=1)
    _kv_table(doc, [("Prepared by", c["assigned_to"]), ("Requested by", c["requested_by"]),
                    ("Reviewed by", "Release Management Office")])
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = foot.add_run("Synthetic demo document for a fictional company.")
    fr.font.size, fr.font.color.rgb, fr.italic = Pt(8), GREY, True
    doc.save(path)


def tr_list_xlsx(c, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "TR list"
    ws["A1"] = f"TR list · {c['change_id']} · {c['short_description']}"
    ws["A1"].font = Font(bold=True, size=13, color="1F2A44")
    ws["A2"] = f"{c['application']} · {c['system']} · release {c['requested_release']}"
    ws["A2"].font = Font(color="6B7280")
    headers = ["Seq", "TR", "Type", "Description", "A target", "P target", "Import after",
               "Change"]
    ws.append([])
    ws.append(headers)
    fill = PatternFill("solid", fgColor="1F2A44")
    thin = Side(style="thin", color="D0D5DD")
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    for t in c["tr_list"]:
        ws.append([t["sequence"], t["tr_id"], t["type"], t["description"], t["a_target"],
                   t["p_target"], t["depends_on_tr"] or "", c["change_id"]])
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for cell in row:
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for col, width in zip("ABCDEFGH", [6, 16, 22, 46, 22, 22, 16, 11]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"
    wb.save(path)


def write_attachments(changes, releases, data_dir):
    """Write the files and record them on each change as ticket attachments."""
    rel = {r["release_id"]: r for r in releases}
    root = Path(data_dir) / "attachments"
    for c in changes:
        c["attachments"] = []
        if not c["release_note"]:
            continue
        folder = root / c["change_id"]
        folder.mkdir(parents=True, exist_ok=True)
        r = rel.get(c["requested_release"], {})
        window = f"{r.get('p_start', '')[:16].replace('T', ' ')} – " \
                 f"{r.get('p_end', '')[11:16]} IST" if r else ""
        rn_name = f"Release_Note_{c['change_id']}.docx"
        tr_name = f"TR_List_{c['change_id']}.xlsx"
        release_note_docx(c, r.get("name", ""), window, folder / rn_name)
        tr_list_xlsx(c, folder / tr_name)
        c["attachments"] = [
            {"name": rn_name, "kind": "Release note",
             "path": f"attachments/{c['change_id']}/{rn_name}"},
            {"name": tr_name, "kind": "TR list",
             "path": f"attachments/{c['change_id']}/{tr_name}"},
        ]
